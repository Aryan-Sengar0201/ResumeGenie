import os
from pathlib import Path
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from docx import Document

SAMPLES_DIR = Path(__file__).resolve().parent.parent / "sample_resumes"
SAMPLES_DIR.mkdir(parents=True, exist_ok=True)

def generate_sample_swe_pdf():
    pdf_path = SAMPLES_DIR / "sample_software_engineer.pdf"
    doc = SimpleDocTemplate(str(pdf_path), pagesize=letter, rightMargin=54, leftMargin=54, topMargin=54, bottomMargin=54)
    styles = getSampleStyleSheet()

    header_style = ParagraphStyle('Head', parent=styles['Normal'], fontName='Helvetica-Bold', fontSize=18, leading=22)
    contact_style = ParagraphStyle('Contact', parent=styles['Normal'], fontName='Helvetica', fontSize=10, leading=14)
    sec_style = ParagraphStyle('Sec', parent=styles['Normal'], fontName='Helvetica-Bold', fontSize=13, leading=17, spaceBefore=10, spaceAfter=4)
    body_style = ParagraphStyle('Body', parent=styles['Normal'], fontName='Helvetica', fontSize=9.5, leading=14)

    story = [
        Paragraph("Alex Mercer", header_style),
        Paragraph("alex.mercer@example.com | +1 (555) 234-5678 | linkedin.com/in/alexmercer-dev | github.com/alexmercer", contact_style),
        Spacer(1, 10),
        
        Paragraph("Professional Summary", sec_style),
        Paragraph("Passionate Full Stack Software Engineer with foundational experience developing scalable web applications using React, TypeScript, Node.js, and Python. Proven track record in optimizing backend REST APIs and containerizing services with Docker.", body_style),
        
        Paragraph("Education", sec_style),
        Paragraph("<b>Bachelor of Science in Computer Science</b> — University of California (2020 – 2024)<br/>GPA: 3.8/4.0. Relevant Coursework: Data Structures, Distributed Systems, Database Management, Cloud Computing.", body_style),
        
        Paragraph("Technical Skills", sec_style),
        Paragraph("<b>Languages:</b> Python, JavaScript, TypeScript, SQL, HTML5, CSS3, C++<br/>"
                  "<b>Frameworks & Libraries:</b> React, Node.js, Express, FastAPI, Tailwind CSS, Redux<br/>"
                  "<b>Databases & Cloud:</b> PostgreSQL, MongoDB, Redis, Docker, AWS, Git, GitHub Actions<br/>"
                  "<b>Soft Skills:</b> Team Leadership, Problem Solving, Agile / Scrum, Technical Communication", body_style),
                  
        Paragraph("Work Experience", sec_style),
        Paragraph("<b>Software Engineering Intern</b> — TechNova Solutions (June 2023 – Sept 2023)<br/>"
                  "• Engineered 12+ RESTful API microservices in Node.js and PostgreSQL, reducing server response times by 32% across 50,000+ daily queries.<br/>"
                  "• Architected a responsive web portal using React and TypeScript, boosting customer onboarding completion rate by 24%.<br/>"
                  "• Automated continuous integration test pipelines using GitHub Actions and Jest, decreasing deployment failures by 40%.", body_style),
                  
        Paragraph("Key Projects", sec_style),
        Paragraph("<b>CloudStream — High-Throughput Media Service</b> | <i>React, FastAPI, PostgreSQL, Docker</i><br/>"
                  "• Developed an asynchronous streaming application capable of processing 1,000+ concurrent audio feeds.<br/>"
                  "• Deployed containerized Docker workloads on AWS EC2, slashing infrastructure hosting costs by $350/month.<br/>"
                  "• Integrated Redis caching layer, decreasing database query latency by 45%.", body_style),
                  
        Paragraph("Certifications & Awards", sec_style),
        Paragraph("• AWS Certified Cloud Practitioner (2023)<br/>• 1st Place Winner — University Annual Hackathon (50+ teams)", body_style)
    ]

    doc.build(story)
    return pdf_path

def generate_sample_data_docx():
    docx_path = SAMPLES_DIR / "sample_data_analyst.docx"
    doc = Document()

    doc.add_heading("Maya Lin", level=1)
    doc.add_paragraph("maya.lin@example.com | (555) 789-0123 | linkedin.com/in/mayalin-data | github.com/mayalin")
    
    doc.add_heading("Professional Summary", level=2)
    doc.add_paragraph("Detail-oriented Data Analyst skilled in Python, SQL, Pandas, Tableau, and Power BI. Experienced in quantitative data modeling, KPI tracking dashboards, and translating complex datasets into actionable business intelligence.")

    doc.add_heading("Education", level=2)
    doc.add_paragraph("Bachelor of Science in Information Systems — State University (2021 – 2025)\nDean's Honor List, Coursework in Business Intelligence, Predictive Analytics, Applied Statistics.")

    doc.add_heading("Technical Skills", level=2)
    # Using a table for skills to demonstrate robust docx table parsing!
    table = doc.add_table(rows=4, cols=2)
    table.rows[0].cells[0].text = "Languages & Tools"
    table.rows[0].cells[1].text = "Python, SQL, R, SQLite, Excel (Advanced)"
    table.rows[1].cells[0].text = "Data Libraries"
    table.rows[1].cells[1].text = "Pandas, NumPy, Scikit-Learn, Matplotlib"
    table.rows[2].cells[0].text = "Visualization & BI"
    table.rows[2].cells[1].text = "Tableau, Power BI, Data Analysis, Git"
    table.rows[3].cells[0].text = "Competencies"
    table.rows[3].cells[1].text = "Problem Solving, Critical Thinking, Presentation"

    doc.add_heading("Projects & Academic Experience", level=2)
    doc.add_paragraph("E-Commerce Customer Retention & Churn Analysis\n"
                      "• Analyzed 150,000+ consumer transaction records using Python and Pandas to identify retention drivers.\n"
                      "• Designed an interactive Tableau dashboard tracking monthly churn rate, increasing marketing campaign efficiency by 22%.\n"
                      "• Implemented SQL queries to automate monthly KPI reporting, saving 15 hours of manual data entry per week.")

    doc.add_paragraph("Healthcare Inventory Demand Forecasting Model\n"
                      "• Formulated predictive regression models in Scikit-Learn, achieving 89% forecast accuracy on historical clinic supply usage.\n"
                      "• Visualized supply chain metrics in Power BI for non-technical stakeholders.")

    doc.add_heading("Certifications", level=2)
    doc.add_paragraph("Google Data Analytics Professional Certificate (2024)")

    doc.save(str(docx_path))
    return docx_path

if __name__ == "__main__":
    generate_sample_swe_pdf()
    generate_sample_data_docx()
    print("Sample resumes generated successfully.")
